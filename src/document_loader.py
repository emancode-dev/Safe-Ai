"""Safe, in-memory extraction for supported upload formats."""

from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from pathlib import PurePath
from typing import BinaryIO

MAX_UPLOAD_BYTES = 10 * 1024 * 1024
SUPPORTED_EXTENSIONS = {
    ".txt", ".md", ".log", ".json", ".yaml", ".yml", ".py", ".js", ".ts",
    ".env", ".example",
}
DOCUMENT_EXTENSIONS = {".pdf", ".docx"}


class UploadError(ValueError):
    """Raised when an upload cannot be safely decoded."""


@dataclass(frozen=True)
class ExtractedDocument:
    filename: str
    file_type: str
    text: str
    character_count: int


def _extension(filename: str) -> str:
    name = PurePath(filename).name.lower()
    if name == ".env.example":
        return ".env.example"
    return PurePath(name).suffix


def _read_bytes(uploaded_file: BinaryIO) -> bytes:
    data = uploaded_file.read(MAX_UPLOAD_BYTES + 1)
    if len(data) > MAX_UPLOAD_BYTES:
        raise UploadError("The upload exceeds the 10 MB size limit.")
    return data


def _extract_pdf(data: bytes) -> str:
    try:
        import fitz
        with fitz.open(stream=data, filetype="pdf") as document:
            return "\n".join(page.get_text() for page in document)
    except Exception as exc:
        raise UploadError("The PDF could not be read safely.") from exc


def _extract_docx(data: bytes) -> str:
    try:
        from docx import Document
        document = Document(BytesIO(data))
        return "\n".join(paragraph.text for paragraph in document.paragraphs)
    except Exception as exc:
        raise UploadError("The DOCX could not be read safely.") from exc


def extract_upload(uploaded_file: BinaryIO) -> ExtractedDocument:
    """Extract readable text without persisting or executing upload content."""
    filename = PurePath(getattr(uploaded_file, "name", "upload")).name
    extension = _extension(filename)
    supported = SUPPORTED_EXTENSIONS | DOCUMENT_EXTENSIONS | {".env.example"}
    if extension not in supported and filename.lower() != ".env.example":
        raise UploadError(
            "Unsupported file type. Use text/config/code files, PDF, or DOCX."
        )

    data = _read_bytes(uploaded_file)
    if extension == ".pdf":
        text = _extract_pdf(data)
        file_type = "PDF"
    elif extension == ".docx":
        text = _extract_docx(data)
        file_type = "DOCX"
    else:
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise UploadError("The file is not valid UTF-8 text.") from exc
        file_type = extension.lstrip(".").upper() or "TEXT"

    if not text.strip():
        raise UploadError("The file contains no readable text.")
    return ExtractedDocument(filename, file_type, text, len(text))
